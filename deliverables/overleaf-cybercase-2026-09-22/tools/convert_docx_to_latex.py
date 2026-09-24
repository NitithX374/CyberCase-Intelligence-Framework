from __future__ import annotations

import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[3]
OUTPUT = Path(__file__).resolve().parents[1]
SOURCES = {
    "3": ROOT / "deliverables/thesis-chapter3-2026-09-21/Chapter3_CyberCase_Methodology_English_AngsanaNew.docx",
    "4": ROOT / "deliverables/thesis-chapter4-2026-09-21/Chapter4_CyberCase_System_Development_English_AngsanaNew.docx",
}


def iter_blocks(document: DocumentType):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def normalized_text(value: str) -> str:
    return " ".join(value.split())


def latex_escape(value: str) -> str:
    replacements = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
        "`": r"\textasciigrave{}",
    }
    return "".join(replacements.get(character, character) for character in value)


def caption_parts(document: DocumentType, prefix: str) -> list[tuple[str, str]]:
    pattern = re.compile(rf"^{prefix}\s+(\d+-\d+)\s+(.+)$")
    captions = []
    for paragraph in document.paragraphs:
        match = pattern.match(normalized_text(paragraph.text))
        if match:
            captions.append((match.group(1), match.group(2)))
    return captions


def image_blob(document: DocumentType, paragraph: Paragraph) -> tuple[str, bytes] | None:
    for element in paragraph._p.iter():
        if element.tag != qn("a:blip"):
            continue
        relationship_id = element.get(qn("r:embed"))
        if relationship_id is None:
            continue
        part = document.part.related_parts[relationship_id]
        return Path(str(part.partname)).name, part.blob
    return None


def write_image(document: DocumentType, paragraph: Paragraph, output: Path, figure_number: str) -> str:
    image = image_blob(document, paragraph)
    if image is None:
        raise ValueError(f"Figure {figure_number} has no embedded image")
    original_name, blob = image
    suffix = Path(original_name).suffix or ".png"
    target = output / f"figure-{figure_number}{suffix}"
    target.write_bytes(blob)
    return target.name


def collapsed_cells(row: object) -> list[str]:
    values: list[str] = []
    for cell in row.cells:
        value = "\n".join(line.strip() for line in cell.text.splitlines() if line.strip())
        if not values or value != values[-1]:
            values.append(value)
    return values


def column_spec(count: int) -> str:
    widths = {
        2: ("0.28\\textwidth", "0.68\\textwidth"),
        3: ("0.20\\textwidth", "0.30\\textwidth", "0.46\\textwidth"),
        4: ("0.15\\textwidth", "0.19\\textwidth", "0.46\\textwidth", "0.14\\textwidth"),
    }.get(count)
    if widths is None:
        widths = tuple("0.9\\textwidth" for _ in range(count))
    return " ".join(
        rf">{{\raggedright\arraybackslash}}p{{{width}}}" for width in widths
    )


def cell_value(value: str) -> str:
    lines = [latex_escape(line.strip()) for line in value.splitlines() if line.strip()]
    return r" \newline ".join(lines)


def render_table(table: Table, caption: tuple[str, str] | None) -> list[str]:
    rows = [collapsed_cells(row) for row in table.rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return []
    count = max(len(row) for row in rows)
    rows = [row + [""] * (count - len(row)) for row in rows]
    caption_text = caption[1] if caption else "Data table"
    label = caption[0] if caption else "table"
    output = [
        r"\begin{longtable}{@{}" + column_spec(count) + r"@{}}",
        rf"\caption{{{latex_escape(caption_text)}}}\label{{tab:{label}}}\\",
        r"\toprule",
        " & ".join(rf"\textbf{{{cell_value(value)}}}" for value in rows[0]) + r" \\",
        r"\midrule",
        r"\endfirsthead",
        " & ".join(rf"\textbf{{{cell_value(value)}}}" for value in rows[0]) + r" \\",
        r"\midrule",
        r"\endhead",
    ]
    for row in rows[1:]:
        output.append(" & ".join(cell_value(value) for value in row) + r" \\")
    output.extend([r"\bottomrule", r"\end{longtable}", ""])
    return output


def render_figure(path: str, caption: tuple[str, str]) -> list[str]:
    number, text = caption
    return [
        r"\begin{figure}[H]",
        r"\centering",
        rf"\includegraphics[width=0.96\textwidth]{{{path}}}",
        rf"\caption{{{latex_escape(text)}}}\label{{fig:{number}}}",
        r"\end{figure}",
        "",
    ]


def render_paragraph(text: str) -> list[str]:
    text = normalized_text(text)
    if not text:
        return []
    section = re.match(r"^(\d+\.\d+(?:\.\d+)?)\s+(.+)$", text)
    if section:
        command = "section" if section.group(1).count(".") == 1 else "subsection"
        return [rf"\{command}{{{latex_escape(section.group(2))}}}", ""]
    numbered = re.match(r"^(\d+)\)\s+(.+)$", text)
    if numbered:
        return [rf"\noindent\textbf{{{numbered.group(1)})}} {latex_escape(numbered.group(2))}", ""]
    return [latex_escape(text), ""]


def render_chapter(chapter_number: str, source: Path) -> None:
    document = Document(str(source))
    figure_captions = caption_parts(document, "Figure")
    table_captions = caption_parts(document, "Table")
    figure_index = 0
    table_index = 0
    blocks = list(iter_blocks(document))
    output: list[str] = []
    chapter_title = next(
        normalized_text(paragraph.text)
        for paragraph in document.paragraphs
        if normalized_text(paragraph.text).startswith("Chapter ")
    )
    title_paragraphs = [normalized_text(paragraph.text) for paragraph in document.paragraphs]
    title_position = title_paragraphs.index(chapter_title)
    full_title = title_paragraphs[title_position + 1]
    output.extend([rf"\chapter{{{latex_escape(full_title)}}}", ""])

    for block in blocks:
        if isinstance(block, Paragraph):
            text = normalized_text(block.text)
            if (
                text.startswith("Chapter ")
                or text == full_title
                or re.match(r"^(Figure|Table)\s+\d+-\d+\s+", text)
            ):
                continue
            if image_blob(document, block) is not None:
                figure_index += 1
                if figure_index > len(figure_captions):
                    raise ValueError(f"Missing caption for figure {chapter_number}-{figure_index}")
                figure_number = figure_captions[figure_index - 1][0]
                image_name = write_image(
                    document,
                    block,
                    OUTPUT / "figures" / f"chapter{chapter_number}",
                    figure_number,
                )
                relative_path = f"figures/chapter{chapter_number}/{image_name}"
                output.extend(render_figure(relative_path, figure_captions[figure_index - 1]))
                continue
            output.extend(render_paragraph(text))
            continue
        table_index += 1
        if table_index > len(table_captions):
            raise ValueError(f"Missing caption for table {chapter_number}-{table_index}")
        output.extend(render_table(block, table_captions[table_index - 1]))

    target = OUTPUT / "chapters" / f"chapter{chapter_number}.tex"
    target.write_text("\n".join(output).rstrip() + "\n", encoding="utf-8")


def main() -> None:
    for chapter_number, source in SOURCES.items():
        render_chapter(chapter_number, source)


if __name__ == "__main__":
    main()
