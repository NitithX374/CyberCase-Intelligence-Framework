"""Build Chapter 4, System Development, as a .docx in the KMUTNB book format.

Every code figure is sliced out of the repository at build time rather than
retyped, so the listing in the book is the code that ran. The commit is
recorded on the first page.

    python docs/thesis/build_chapter4.py
"""

from __future__ import annotations

import ast
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

REPO = Path(__file__).resolve().parents[2]
OUTPUT = REPO / "docs" / "thesis" / "Chapter4_System_Development.docx"

BODY_FONT = "TH SarabunPSK"
CODE_FONT = "Consolas"
BODY_SIZE = Pt(16)
CODE_SIZE = Pt(10)


def commit() -> str:
    try:
        return subprocess.check_output(
            ["git", "rev-parse", "--short", "HEAD"], cwd=REPO, text=True
        ).strip()
    except Exception:
        return "unknown"


def source_of(relative: str, name: str | None = None) -> str:
    """A file, or one class or function inside it, exactly as it is on disk."""

    path = REPO / relative
    text = path.read_text(encoding="utf-8")
    if name is None:
        return text.rstrip()
    tree = ast.parse(text)
    lines = text.splitlines()
    for node in ast.walk(tree):
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
            if node.name != name:
                continue
            start = min([node.lineno] + [d.lineno for d in node.decorator_list]) - 1
            return "\n".join(lines[start : node.end_lineno]).rstrip()
    raise KeyError(f"{name} not found in {relative}")


class Chapter:
    def __init__(self) -> None:
        self.document = Document()
        self.figure = 0
        self.table_number = 0
        self._styles()

    def _styles(self) -> None:
        normal = self.document.styles["Normal"]
        normal.font.name = BODY_FONT
        normal.font.size = BODY_SIZE
        normal.element.rPr.rFonts.set(qn("w:cs"), BODY_FONT)
        for name, size in (("Heading 1", 20), ("Heading 2", 16), ("Heading 3", 16)):
            style = self.document.styles[name]
            style.font.name = BODY_FONT
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.element.rPr.rFonts.set(qn("w:cs"), BODY_FONT)

    def heading(self, text: str, level: int = 2) -> None:
        self.document.add_heading(text, level=level)

    def body(self, text: str) -> None:
        paragraph = self.document.add_paragraph(text)
        paragraph.paragraph_format.first_line_indent = Pt(36)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def plain(self, text: str, bold: bool = False) -> None:
        paragraph = self.document.add_paragraph()
        run = paragraph.add_run(text)
        run.bold = bold
        paragraph.paragraph_format.space_after = Pt(4)

    def bullets(self, items: list[str]) -> None:
        for item in items:
            paragraph = self.document.add_paragraph(item, style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(2)

    @staticmethod
    def _span(code: str, spec) -> str:
        """Turn an anchor, or a pair of them, into a printed line range.

        Hand-written line numbers go stale the moment the code moves. An anchor
        is a distinctive fragment of the line it names, so the range printed in
        the book is always the range of the listing above it. A missing anchor
        fails the build rather than printing a wrong number.
        """

        if isinstance(spec, str):
            return spec
        lines = code.splitlines()

        def find(fragment: str) -> int:
            # A leading ~ means the last occurrence, for a line such as
            # "return None" that appears several times in one listing.
            last = fragment.startswith("~")
            fragment = fragment.removeprefix("~")
            # An exact line wins over a line that merely contains the text, so
            # "return evaluate_x" does not match "return evaluate_x_encoder".
            for test in (lambda line: line.strip() == fragment.strip(), lambda line: fragment in line):
                hits = [index for index, line in enumerate(lines, 1) if test(line)]
                if hits:
                    return hits[-1] if last else hits[0]
            raise KeyError(f"anchor not found in listing: {fragment!r}")

        start = find(spec[0])
        if len(spec) == 1:
            return f"Line {start}"
        end = find(spec[1])
        return f"Line {start}" if start == end else f"Lines {start}-{end}"

    def code_figure(self, code: str, caption: str, notes: list) -> None:
        """A numbered code listing followed by its line-range commentary."""

        self.figure += 1
        numbered = "\n".join(
            f"{index:>3}  {line}" for index, line in enumerate(code.splitlines(), 1)
        )
        paragraph = self.document.add_paragraph()
        run = paragraph.add_run(numbered)
        run.font.name = CODE_FONT
        run.font.size = CODE_SIZE
        run.element.rPr.rFonts.set(qn("w:cs"), CODE_FONT)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.left_indent = Pt(18)

        label = self.document.add_paragraph()
        label_run = label.add_run(f"Figure 4-{self.figure}  {caption}")
        label_run.bold = True
        label.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label.paragraph_format.space_after = Pt(6)

        for spec, explanation in notes:
            note = self.document.add_paragraph()
            note.paragraph_format.left_indent = Pt(18)
            note.paragraph_format.space_after = Pt(0)
            head = note.add_run(f"{self._span(code, spec)}\t")
            head.bold = True
            note.add_run(explanation)
        self.document.add_paragraph()

    def table(self, caption: str, header: list[str], rows: list[list[str]]) -> None:
        self.table_number += 1
        label = self.document.add_paragraph()
        run = label.add_run(f"Table 4-{self.table_number}  {caption}")
        run.bold = True
        label.alignment = WD_ALIGN_PARAGRAPH.CENTER

        table = self.document.add_table(rows=1, cols=len(header))
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, header):
            cell.text = ""
            run = cell.paragraphs[0].add_run(text)
            run.bold = True
            run.font.size = Pt(14)
            run.font.name = BODY_FONT
        for values in rows:
            cells = table.add_row().cells
            for cell, text in zip(cells, values):
                cell.text = ""
                run = cell.paragraphs[0].add_run(text)
                run.font.size = Pt(14)
                run.font.name = BODY_FONT
        self.document.add_paragraph()

    def save(self) -> None:
        OUTPUT.parent.mkdir(parents=True, exist_ok=True)
        self.document.save(OUTPUT)
